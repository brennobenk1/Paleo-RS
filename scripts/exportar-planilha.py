#!/usr/bin/env python3
"""
Paleo-RS — scripts/exportar-planilha.py

js/dados.js  ->  planilha.xlsx

Artefato gerado. A planilha existe para quem prefere trabalhar em Excel,
nunca como entrada: no Paleo-SC o fluxo era planilha -> JSON e as cópias
divergiram até o site mostrar 158 registros enquanto data/ tinha 97.

Requer openpyxl:  pip install openpyxl
"""

import sys
from pathlib import Path

from ler_dados import carregar, RAIZ

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl não instalado. Rode: pip install openpyxl")

COLUNAS = ["id", "taxon", "categoria", "periodo", "idade_ma", "formacao", "bacia",
           "municipio", "local_coleta", "site", "lat", "lon", "coord_precisao",
           "numero_catalogo", "armazenamento", "unidade_pesquisa", "descritor",
           "tipo_fonte", "doi", "observacoes", "citacao_abnt"]


def aba(wb, titulo, colunas, linhas):
    ws = wb.create_sheet(titulo)
    ws.append(colunas)
    for c in range(1, len(colunas) + 1):
        cel = ws.cell(row=1, column=c)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="143D28")   # verde do cabeçalho
        cel.alignment = Alignment(vertical="center")
    for linha in linhas:
        ws.append(linha)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c, nome in enumerate(colunas, start=1):
        largura = 60 if nome in ("observacoes", "citacao_abnt", "local_coleta") else \
                  34 if nome in ("taxon", "armazenamento", "descritor", "formacao", "bacia") else 16
        ws.column_dimensions[get_column_letter(c)].width = largura
    return ws


def main():
    d = carregar()
    wb = Workbook()
    wb.remove(wb.active)

    leia = wb.create_sheet("LEIA-ME")
    for i, texto in enumerate([
        f"Paleo-RS — versão {d['VERSAO_BANCO']} ({d['DATA_VERSAO']})",
        "",
        "Esta planilha é um ARTEFATO GERADO a partir de js/dados.js.",
        "Alterações feitas aqui não voltam para o banco e serão perdidas",
        "na próxima exportação. Para corrigir um registro, edite js/dados.js",
        "e rode scripts/validar.py.",
        "",
        "Dados sob CC BY 4.0. Citar o banco não substitui citar a fonte",
        "primária de cada registro.",
    ], start=1):
        leia.cell(row=i, column=1, value=texto)
    leia.column_dimensions["A"].width = 76
    leia.cell(row=1, column=1).font = Font(bold=True, size=13)

    aba(wb, "Registros", COLUNAS,
        [[r.get(c, "") for c in COLUNAS] for r in d["registros"]])
    aba(wb, "Sítios", ["nome", "municipio", "lat", "lon", "count", "coord_precisao"],
        [[s.get(k, "") for k in ("nome", "municipio", "lat", "lon", "count", "coord_precisao")]
         for s in d["sitios"]])
    aba(wb, "Períodos", ["chave", "nome", "ordem", "inicio_ma", "fim_ma", "cor", "total_registros"],
        [[p.get(k, "") for k in ("chave", "nome", "ordem", "inicio_ma", "fim_ma", "cor", "total_registros")]
         for p in d["periodos"]])
    aba(wb, "Biozonas", ["chave", "nome", "ordem", "idade", "total_registros"],
        [[z.get(k, "") for k in ("chave", "nome", "ordem", "idade", "total_registros")]
         for z in d["biozonas"]])
    aba(wb, "Instituições", ["sigla", "nome", "cidade", "site", "acervo"],
        [[i.get(k, "") for k in ("sigla", "nome", "cidade", "site", "acervo")]
         for i in d["instituicoes"]])

    destino = RAIZ / "planilha.xlsx"
    wb.save(destino)
    print(f"{destino.relative_to(RAIZ)} — {len(d['registros'])} registros")


if __name__ == "__main__":
    main()
